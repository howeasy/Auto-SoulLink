#!/usr/bin/env python3
"""tools/fetch_rr_sources.py — fetch/verify the pinned upstream sources for
the Radical Red data generators (tools/gen_rr_*.py) against
data/gen3_rr_sources.lock.json. See docs/gen3_requirements.md row F-7.

Every gen_rr_*.py calls `cached_source(name)` from this module to get pinned
bytes for one lock entry: a cache hit whose sha256 matches the pin is
returned as-is; otherwise the pinned `fetch_url` is fetched and the sha256
is verified before it's trusted or cached. A mismatch always raises -- never
a silent pass -- because the pin exists specifically to catch upstream
drift (two of these sources are dead repos preserved only in the Wayback
Machine; a third is a community Google Sheet with no immutable revision).

Usage:
    python tools/fetch_rr_sources.py            # populate $SLINK_RR_SRC_CACHE
    python tools/fetch_rr_sources.py --check     # verify the existing cache only (no network)
"""

from __future__ import annotations

import gzip
import hashlib
import json
import os
import sys
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
LOCK_PATH = ROOT / "data" / "gen3_rr_sources.lock.json"


def cache_dir() -> Path:
    return Path(os.environ.get("SLINK_RR_SRC_CACHE", ROOT / "data" / ".rr_src_cache"))


def load_lock() -> dict:
    return json.loads(LOCK_PATH.read_text(encoding="utf-8"))


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _fetch(url: str) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": "slink-rr-fetch/1"})
    with urllib.request.urlopen(req, timeout=30) as resp:
        raw = resp.read()
    # A Wayback Machine "id_" replay serves the exact bytes GitHub sent at
    # capture time, gzip Content-Encoding included -- decode it the same way
    # a browser/curl --compressed would, whether or not the header survived.
    if raw[:2] == b"\x1f\x8b":
        raw = gzip.decompress(raw)
    return raw


def _xlsx_content_sha256(data: bytes) -> str:
    """Hash an xlsx's actual cell content, not its zip bytes.

    Google's xlsx exporter rewrites xl/sharedStrings.xml and every
    xl/worksheets/*.xml non-deterministically on each export (confirmed by
    exporting the same unedited sheet twice and diffing) -- a raw sha256
    would never match twice. This hashes openpyxl's (sheet, row, col,
    repr(value)) tuples with data_only=True instead, which reproducible
    across exports. See the rr_priority_trainers_sheet_xlsx lock note.
    """
    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True)
    rows = [
        (ws.title, cell.row, cell.column, repr(cell.value))
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    ]
    return _sha256(json.dumps(rows, ensure_ascii=False).encode("utf-8"))


def _verify(entry: dict, data: bytes) -> tuple[bool, str, str]:
    """Return (ok, expected, got) for whichever pin field this entry uses."""
    if entry.get("kind") == "google_sheets_xlsx":
        expected = entry["content_sha256"]
        got = _xlsx_content_sha256(data)
    else:
        expected = entry["sha256"]
        got = _sha256(data)
    return got == expected, expected, got


def _cache_path(name: str, entry: dict) -> Path:
    # openpyxl refuses to open a zip unless the filename ends in .xlsx (it
    # sniffs the extension, not the content) -- give that one entry a real
    # extension on disk; every other source is read as opaque bytes/text.
    suffix = ".xlsx" if entry.get("kind") == "google_sheets_xlsx" else ""
    return cache_dir() / f"{name}{suffix}"


def cached_source(name: str, *, allow_fetch: bool = True) -> bytes:
    """Return pinned, hash-verified bytes for lock entry `name`.

    Prefers the on-disk cache; fetches from the pinned URL and populates the
    cache when the entry is missing/stale and `allow_fetch` is true.
    """
    lock = load_lock()
    try:
        entry = lock["sources"][name]
    except KeyError:
        raise KeyError(f"{name!r} is not declared in {LOCK_PATH}") from None
    path = _cache_path(name, entry)
    if path.exists():
        data = path.read_bytes()
        ok, expected, got = _verify(entry, data)
        if ok:
            return data
        if not allow_fetch:
            raise ValueError(
                f"{name}: cached file hash mismatch (expected {expected}, got {got})"
            )
    if not allow_fetch:
        raise FileNotFoundError(
            f"{name}: not cached at {path} (run tools/fetch_rr_sources.py first, "
            f"or set SLINK_RR_SRC_CACHE)"
        )
    data = _fetch(entry["fetch_url"])
    ok, expected, got = _verify(entry, data)
    if not ok:
        raise ValueError(
            f"{name}: fetched hash {got} != pinned {expected} "
            f"(source drifted -- re-pin deliberately in {LOCK_PATH}, don't just accept this)"
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return data


def cached_source_path(name: str, *, allow_fetch: bool = True) -> Path:
    """Like cached_source(), but returns the on-disk path (for tools, like
    openpyxl, that need a real file rather than bytes)."""
    cached_source(name, allow_fetch=allow_fetch)
    return _cache_path(name, load_lock()["sources"][name])


def main(argv: list[str] | None = None) -> int:
    argv = sys.argv[1:] if argv is None else argv
    check_only = "--check" in argv
    lock = load_lock()
    ok = True
    for name in lock["sources"]:
        try:
            data = cached_source(name, allow_fetch=not check_only)
        except Exception as exc:
            print(f"FAIL {name}: {exc}")
            ok = False
            continue
        print(f"OK   {name}: {len(data)} bytes")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
